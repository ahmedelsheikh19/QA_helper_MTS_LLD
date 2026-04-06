"""
Shared fixtures for building in-memory Excel files used across the test suite.
All helpers return a Path to a saved .xlsx file inside pytest's tmp_path.
"""
import pytest
import openpyxl


# ── Generic builder ────────────────────────────────────────────────────────

def make_xlsx(tmp_path, filename, sheets):
    """
    Create an xlsx file at tmp_path/filename.

    sheets: dict of { sheet_name: [[row], [row], ...] }
            The first sheet dict entry becomes the active sheet.
    """
    wb = openpyxl.Workbook()
    first = True
    for name, rows in sheets.items():
        if first:
            ws = wb.active
            ws.title = name
            first = False
        else:
            ws = wb.create_sheet(name)
        for row in rows:
            ws.append(row)
    path = tmp_path / filename
    wb.save(path)
    return path


# ── Requirements file helpers ──────────────────────────────────────────────

@pytest.fixture
def req_file_id_col(tmp_path):
    """Requirements file with 'ID' column."""
    return make_xlsx(tmp_path, "reqs.xlsx", {
        "Sheet1": [
            ["ID", "Description"],
            ["REQ-001", "First requirement"],
            ["REQ-002", "Second requirement"],
            ["REQ-003", "Third requirement"],
        ]
    })


@pytest.fixture
def req_file_object_identifier_col(tmp_path):
    """Requirements file with 'Object Identifier' column."""
    return make_xlsx(tmp_path, "reqs_oi.xlsx", {
        "Sheet1": [
            ["Object Identifier", "Description"],
            ["REQ-001", "First requirement"],
            ["REQ-002", "Second requirement"],
        ]
    })


@pytest.fixture
def req_file_no_id_col(tmp_path):
    """Requirements file missing any recognised ID column."""
    return make_xlsx(tmp_path, "reqs_no_id.xlsx", {
        "Sheet1": [
            ["Name", "Description"],
            ["REQ-001", "First requirement"],
        ]
    })


@pytest.fixture
def req_file_empty_rows(tmp_path):
    """Requirements file with header only, no data rows."""
    return make_xlsx(tmp_path, "reqs_empty.xlsx", {
        "Sheet1": [["ID", "Description"]]
    })


@pytest.fixture
def req_file_duplicates(tmp_path):
    """Requirements file with duplicate IDs."""
    return make_xlsx(tmp_path, "reqs_dup.xlsx", {
        "Sheet1": [
            ["ID"],
            ["REQ-001"],
            ["REQ-001"],
            ["REQ-002"],
        ]
    })


@pytest.fixture
def req_file_whitespace(tmp_path):
    """Requirements file with IDs that have leading/trailing whitespace."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["ID"])
    ws.append(["  REQ-001  "])
    ws.append([" REQ-002"])
    path = tmp_path / "reqs_ws.xlsx"
    wb.save(path)
    return path


# ── LLD file helpers ───────────────────────────────────────────────────────

@pytest.fixture
def lld_file_single_reqs(tmp_path):
    """LLD file where each REQ cell contains a single requirement."""
    return make_xlsx(tmp_path, "lld.xlsx", {
        "Sheet1": [
            ["Component", "REQ"],
            ["CompA", "REQ-001"],
            ["CompB", "REQ-002"],
        ]
    })


@pytest.fixture
def lld_file_multi_reqs(tmp_path):
    """LLD file with multiple reqs per cell (comma/space/newline separated)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Component", "REQ"])
    ws.cell(row=2, column=1).value = "CompA"
    ws.cell(row=2, column=2).value = "REQ-001, REQ-002"
    ws.cell(row=3, column=1).value = "CompB"
    ws.cell(row=3, column=2).value = "REQ-003\nREQ-004"
    ws.cell(row=4, column=1).value = "CompC"
    ws.cell(row=4, column=2).value = "REQ-005 REQ-006"
    path = tmp_path / "lld_multi.xlsx"
    wb.save(path)
    return path


@pytest.fixture
def lld_file_no_req_col(tmp_path):
    """LLD file missing a REQ column."""
    return make_xlsx(tmp_path, "lld_no_req.xlsx", {
        "Sheet1": [
            ["Component", "Description"],
            ["CompA", "Does something"],
        ]
    })


@pytest.fixture
def lld_file_empty_rows(tmp_path):
    """LLD file with REQ column but no data rows."""
    return make_xlsx(tmp_path, "lld_empty.xlsx", {
        "Sheet1": [["Component", "REQ"]]
    })


# ── UT file helpers ────────────────────────────────────────────────────────

@pytest.fixture
def ut_file_valid(tmp_path):
    """UT file with properly formatted A_REQUIREMENT_LABEL entries."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "TestSuite"
    ws.cell(row=1, column=1).value = "A_REQUIREMENT_LABEL"
    ws.cell(row=1, column=2).value = "REQ-001,REQ-002,"
    ws.cell(row=2, column=1).value = "A_REQUIREMENT_LABEL"
    ws.cell(row=2, column=2).value = "REQ-003,"
    path = tmp_path / "ut_valid.xlsx"
    wb.save(path)
    return path


@pytest.fixture
def ut_file_format_errors(tmp_path):
    """UT file with various formatting violations."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "TestSuite"
    # Has spaces
    ws.cell(row=1, column=1).value = "A_REQUIREMENT_LABEL"
    ws.cell(row=1, column=2).value = "REQ-001 REQ-002"
    # Has newline
    ws.cell(row=2, column=1).value = "A_REQUIREMENT_LABEL"
    ws.cell(row=2, column=2).value = "REQ-003\nREQ-004"
    # No trailing comma
    ws.cell(row=3, column=1).value = "A_REQUIREMENT_LABEL"
    ws.cell(row=3, column=2).value = "REQ-005"
    # All three violations
    ws.cell(row=4, column=1).value = "A_REQUIREMENT_LABEL"
    ws.cell(row=4, column=2).value = "REQ-006 REQ-007\n"
    path = tmp_path / "ut_errors.xlsx"
    wb.save(path)
    return path


@pytest.fixture
def ut_file_multi_sheet(tmp_path):
    """UT file with A_REQUIREMENT_LABEL spread across multiple sheets."""
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Suite1"
    ws1.cell(row=1, column=1).value = "A_REQUIREMENT_LABEL"
    ws1.cell(row=1, column=2).value = "REQ-001,"
    ws2 = wb.create_sheet("Suite2")
    ws2.cell(row=1, column=1).value = "A_REQUIREMENT_LABEL"
    ws2.cell(row=1, column=2).value = "REQ-002,"
    path = tmp_path / "ut_multisheet.xlsx"
    wb.save(path)
    return path


@pytest.fixture
def ut_file_no_labels(tmp_path):
    """UT file with no A_REQUIREMENT_LABEL cells at all."""
    return make_xlsx(tmp_path, "ut_no_labels.xlsx", {
        "Sheet1": [
            ["TestName", "Status"],
            ["Test_A", "PASS"],
        ]
    })
