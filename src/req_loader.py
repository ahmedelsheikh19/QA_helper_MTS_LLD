"""
req_loader.py — Shared module for loading and filtering requirements from Excel files.

Used by both generate_traceabilityLLD.py and generate_ut_traceability.py,
as well as the GUI. Keeps all requirement-loading logic in one place.
"""

import openpyxl
import logging

# ── Filterable columns and their known values ────────────────────────────────
# Add new columns here as needed — the rest of the code adapts automatically.
FILTER_COLUMNS = {
    "aObject Status": [
        "n/a",
        "Undefined",
        "in Work",
        "in Review",
        "Accepted",
        "Cancelled",
    ],
    "aObject Type": [
        "Undefined",
        "Heading",
        "Doc Info",
        "Information",
        "Req functional",
        "Req non functional",
    ],
}


# ── Core loader ──────────────────────────────────────────────────────────────

def load_requirements(file_path, filters=None, filter_logic="AND"):
    """
    Read a Requirements Excel file and return matching requirement IDs.

    Args:
        file_path:    Path to the .xlsx file.
        filters:      Optional list of filter dicts, each with:
                        {"column": str, "values": [str, ...]}
                      When None or empty, all requirements are returned.
        filter_logic: "AND" (row must match ALL filters) or
                      "OR"  (row must match ANY filter).

    Returns:
        set[str] of requirement IDs.
    """
    logging.info(f"Loading Requirements file: {file_path}")
    wb = openpyxl.load_workbook(file_path, data_only=True)
    sheet = wb.active

    # ── Build a header map: column_name -> column_index ──────────────
    headers = {}
    for col_idx, cell in enumerate(sheet[1], start=1):
        if cell.value:
            headers[str(cell.value).strip()] = col_idx

    # ── Locate the ID column ─────────────────────────────────────────
    id_col = None
    for name in ("ID", "Object Identifier"):
        if name in headers:
            id_col = headers[name]
            break
        # Case-insensitive fallback
        for header_name, idx in headers.items():
            if header_name.lower() == name.lower():
                id_col = idx
                break
        if id_col:
            break

    if not id_col:
        logging.error("Could not find an 'ID' or 'Object Identifier' column.")
        return set()

    logging.info(f"Found ID column at index {id_col}.")

    # ── Resolve filter columns ───────────────────────────────────────
    filter_cols = []
    if filters:
        for f in filters:
            col_name = f["column"]
            if col_name not in headers:
                logging.warning(f"Filter column '{col_name}' not found in file — skipping.")
                continue
            filter_cols.append({
                "name": col_name,
                "index": headers[col_name],
                "values": {v.strip().lower() for v in f["values"]},
                "seen": set(),
            })

        if not filter_cols:
            logging.warning("No valid filter columns found. Returning all requirements.")

    # ── Extract IDs with optional filtering ──────────────────────────
    req_ids = set()

    for row in sheet.iter_rows(min_row=2):
        cell_id = row[id_col - 1].value
        if not cell_id:
            continue

        if filter_cols:
            matches = []
            for fc in filter_cols:
                cell_val = row[fc["index"] - 1].value
                cell_str = str(cell_val).strip().lower() if cell_val else ""
                fc["seen"].add(cell_str)
                matches.append(cell_str in fc["values"])

            if filter_logic.upper() == "AND" and not all(matches):
                continue
            if filter_logic.upper() == "OR" and not any(matches):
                continue

        req_ids.add(str(cell_id).strip())

    # ── Warn about requested filter values that don't exist in the data ──
    for fc in filter_cols:
        unmatched = fc["values"] - fc["seen"]
        if unmatched:
            available = sorted(v for v in fc["seen"] if v)
            logging.warning(
                f"Filter '{fc['name']}': value(s) not found in column → "
                f"{sorted(unmatched)}. Available values: {available}"
            )

    logging.info(f"Extracted {len(req_ids)} requirements (filter_logic={filter_logic}, "
                 f"active_filters={len(filter_cols)}).")
    return req_ids


# ── Schema discovery (used by the GUI) ───────────────────────────────────────

def scan_filterable_columns(file_path):
    """
    Read an Excel file and return {column_name: sorted_unique_values_list}.

    Used by the GUI to populate filter dropdowns dynamically. Excludes the
    ID/Object Identifier column. Drops columns with no non-empty values.
    For duplicate header names, the rightmost column wins (matches the
    behaviour of load_requirements()).
    """
    logging.info(f"Scanning filterable columns in: {file_path}")
    wb = openpyxl.load_workbook(file_path, data_only=True, read_only=True)
    try:
        sheet = wb.active
        rows_iter = sheet.iter_rows(values_only=True)
        try:
            header_row = next(rows_iter)
        except StopIteration:
            return {}

        skip = {"id", "object identifier"}
        header_map = {}  # name -> 0-based column index
        for idx, val in enumerate(header_row):
            if val is None:
                continue
            name = str(val).strip()
            if not name or name.lower() in skip:
                continue
            header_map[name] = idx  # later duplicates overwrite earlier ones

        if not header_map:
            return {}

        unique = {name: set() for name in header_map}
        for row in rows_iter:
            for name, idx in header_map.items():
                if idx >= len(row):
                    continue
                cell_val = row[idx]
                if cell_val is None:
                    continue
                s = str(cell_val).strip()
                if s:
                    unique[name].add(s)

        return {name: sorted(vals) for name, vals in unique.items() if vals}
    finally:
        wb.close()


# ── CLI helpers ──────────────────────────────────────────────────────────────

def add_filter_args(parser):
    """
    Add --filter and --filter-logic arguments to any argparse parser.

    Usage in scripts:
        add_filter_args(parser)
        args = parser.parse_args()
        filters = parse_filters(args.filter)
    """
    parser.add_argument(
        "--filter",
        action="append",
        metavar="COLUMN=VAL1,VAL2",
        help=(
            "Filter requirements by column value. "
            "Repeatable for multiple filters. "
            "Example: --filter \"aObject Status=Accepted,in Review\""
        ),
    )
    parser.add_argument(
        "--filter-logic",
        choices=["AND", "OR"],
        default="AND",
        help="How to combine multiple filters: AND (all match) or OR (any match). Default: AND.",
    )


def parse_filters(filter_args):
    """
    Parse a list of 'COLUMN=VAL1,VAL2' strings into filter dicts.

    Args:
        filter_args: List of strings from argparse (or None).

    Returns:
        List of {"column": str, "values": [str, ...]} dicts.
    """
    if not filter_args:
        return []

    filters = []
    for raw in filter_args:
        if "=" not in raw:
            logging.warning(f"Invalid filter format '{raw}' — expected COLUMN=VAL1,VAL2. Skipping.")
            continue

        col, _, vals = raw.partition("=")
        col = col.strip()
        values = [v.strip() for v in vals.split(",") if v.strip()]

        if not col or not values:
            logging.warning(f"Empty column or values in filter '{raw}'. Skipping.")
            continue

        filters.append({"column": col, "values": values})

    return filters
