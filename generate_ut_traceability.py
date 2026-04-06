import openpyxl
import argparse
import logging
import re
import os

# Configure logging
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(levelname)s - %(message)s'
)

def extract_reqs_from_req_file(file_path):
    """
    Reads the Requirements Excel file and extracts IDs.
    Looks for a column named 'ID' or 'Object Identifier'.
    """
    logging.info(f"Loading Requirements file: {file_path}")
    wb = openpyxl.load_workbook(file_path, data_only=True)
    sheet = wb.active
    
    req_ids = set()
    id_col_idx = None
    
    # Find the column index for ID
    for col_idx, cell in enumerate(sheet[1], start=1):
        if cell.value and str(cell.value).strip().lower() in ['id', 'object identifier']:
            id_col_idx = col_idx
            break
            
    if not id_col_idx:
        logging.error("Could not find an 'ID' or 'Object Identifier' column in the Requirements file.")
        return req_ids
        
    logging.info(f"Found ID column at index {id_col_idx}. Extracting requirements...")
    
    for row in sheet.iter_rows(min_row=2, min_col=id_col_idx, max_col=id_col_idx):
        val = row[0].value
        if val:
            req_ids.add(str(val).strip())
            
    logging.info(f"Extracted {len(req_ids)} unique requirements from Requirements file.")
    return req_ids

def extract_reqs_from_ut_files(file_paths):
    """
    Reads multiple UT Excel files, iterates through all sheets, looks for 
    'A_REQUIREMENT_LABEL', and extracts mapped REQs from the adjacent cell.
    Also validates the strict formatting rules.
    """
    ut_reqs = set()
    format_errors = []

    for file_path in file_paths:
        logging.info(f"Processing UT file: {file_path}")
        try:
            wb = openpyxl.load_workbook(file_path, data_only=True)
        except Exception as e:
            logging.error(f"Failed to load {file_path}: {e}")
            continue
            
        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            logging.info(f"  Scanning sheet: {sheet_name}")
            
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value and str(cell.value).strip() == "A_REQUIREMENT_LABEL":
                        # Get the cell immediately to the right
                        target_cell = sheet.cell(row=cell.row, column=cell.column + 1)
                        val = target_cell.value
                        
                        if val:
                            val_str = str(val)
                            
                            # Validation Checks
                            has_spaces = " " in val_str
                            has_newlines = "\n" in val_str or "\r" in val_str
                            ends_with_comma = val_str.endswith(",")
                            
                            if has_spaces or has_newlines or not ends_with_comma:
                                error_reason = []
                                if has_spaces: error_reason.append("Contains spaces")
                                if has_newlines: error_reason.append("Contains new lines")
                                if not ends_with_comma: error_reason.append("Does not end with a comma")
                                
                                format_errors.append({
                                    "file": os.path.basename(file_path),
                                    "sheet": sheet_name,
                                    "cell": target_cell.coordinate,
                                    "value": val_str,
                                    "reason": " | ".join(error_reason)
                                })
                            
                            # Extract requirements (cleaning up just in case, to still process them)
                            # Split by comma and ignore empty strings (like the one after the last comma)
                            split_reqs = [r.strip() for r in val_str.replace('\n', '').split(',')]
                            for req in split_reqs:
                                if req:
                                    ut_reqs.add(req)
                                    
    logging.info(f"Extracted {len(ut_reqs)} unique requirements from all UT files.")
    logging.info(f"Found {len(format_errors)} formatting errors.")
    return ut_reqs, format_errors

def generate_html_report(req_ids, ut_reqs, format_errors, output_file="ut_traceability_report.html"):
    """
    Generates a modern, interactive HTML report showing the traceability gaps, coverages, and formatting errors.
    """
    logging.info("Generating HTML Report...")
    
    covered_in_ut = req_ids.intersection(ut_reqs)
    missing_in_ut = req_ids - ut_reqs
    missing_in_reqs = ut_reqs - req_ids
    
    # Generate Format Errors HTML list
    format_errors_html = ""
    if format_errors:
        for err in format_errors:
            format_errors_html += f"<li><strong>{err['file']}</strong> -> Sheet: <em>{err['sheet']}</em> -> Cell: <strong>{err['cell']}</strong><br>"
            format_errors_html += f"<span style='color:#e74c3c;'>Error: {err['reason']}</span><br>"
            format_errors_html += f"<div style='background:#eee; padding:5px; margin-top:5px; font-family:monospace;'>{err['value']}</div></li>"
    else:
        format_errors_html = "<li>No formatting errors found! 🎉 All requirements are perfectly formatted.</li>"

    html_content = f"""
    <!DOCTYPE html>
    <html lang="en">
    <head>
        <meta charset="UTF-8">
        <meta name="viewport" content="width=device-width, initial-scale=1.0">
        <title>UT Traceability Matrix Report</title>
        <style>
            body {{ font-family: 'Segoe UI', Tahoma, Geneva, Verdana, sans-serif; background-color: #f4f7f6; margin: 0; padding: 20px; color: #333; }}
            .container {{ max-width: 1000px; margin: auto; background: #fff; padding: 30px; border-radius: 8px; box-shadow: 0 4px 8px rgba(0,0,0,0.1); }}
            h1 {{ text-align: center; color: #2c3e50; }}
            .summary {{ display: flex; justify-content: space-between; margin: 20px 0; gap: 15px; }}
            .card {{ background: #ecf0f1; padding: 20px; border-radius: 8px; text-align: center; flex: 1; box-shadow: 0 2px 4px rgba(0,0,0,0.05); cursor: pointer; transition: 0.3s; }}
            .card:hover {{ background: #bdc3c7; transform: translateY(-3px); }}
            .card h3 {{ margin: 0; font-size: 28px; }}
            .card p {{ margin: 5px 0 0; color: #555; font-weight: bold; font-size: 14px; }}
            
            /* Specific Colors */
            .card.success h3 {{ color: #27ae60; }}
            .card.warning h3 {{ color: #e67e22; }}
            .card.danger h3 {{ color: #e74c3c; }}
            .card.error h3 {{ color: #8e44ad; }}

            .section {{ display: none; margin-top: 30px; animation: fadeIn 0.4s; }}
            .section.active {{ display: block; }}
            @keyframes fadeIn {{ from {{ opacity: 0; }} to {{ opacity: 1; }} }}
            
            h2 {{ color: #34495e; border-bottom: 2px solid #ecf0f1; padding-bottom: 5px; }}
            ul {{ list-style-type: none; padding: 0; max-height: 500px; overflow-y: auto; background: #fafafa; border: 1px solid #ddd; border-radius: 4px; }}
            li {{ padding: 12px; border-bottom: 1px solid #eee; }}
            li:last-child {{ border-bottom: none; }}
        </style>
        <script>
            function toggleSection(sectionId) {{
                document.querySelectorAll('.section').forEach(sec => sec.classList.remove('active'));
                document.getElementById(sectionId).classList.add('active');
            }}
        </script>
    </head>
    <body>
        <div class="container">
            <h1>Unit Test (UT) Traceability Report</h1>
            
            <div class="summary">
                <div class="card success" onclick="toggleSection('covered')">
                    <h3>{len(covered_in_ut)}</h3>
                    <p>Covered in UT</p>
                </div>
                <div class="card danger" onclick="toggleSection('missing-ut')">
                    <h3>{len(missing_in_ut)}</h3>
                    <p>Missing in UT</p>
                </div>
                <div class="card warning" onclick="toggleSection('missing-reqs')">
                    <h3>{len(missing_in_reqs)}</h3>
                    <p>Missing in Reqs File</p>
                </div>
                <div class="card error" onclick="toggleSection('format-errors')">
                    <h3>{len(format_errors)}</h3>
                    <p>Format Violations</p>
                </div>
            </div>

            <div id="covered" class="section active">
                <h2>✅ Covered Requirements (Exist in Both)</h2>
                <ul>
                    {''.join(f'<li>{req}</li>' for req in sorted(covered_in_ut))}
                </ul>
            </div>

            <div id="missing-ut" class="section">
                <h2>❌ Missing in UT (Defined in Reqs, lacking UT coverage)</h2>
                <ul>
                    {''.join(f'<li>{req}</li>' for req in sorted(missing_in_ut))}
                </ul>
            </div>

            <div id="missing-reqs" class="section">
                <h2>⚠️ Unmapped in UT (Exist in UT files, not in Reqs file)</h2>
                <ul>
                    {''.join(f'<li>{req}</li>' for req in sorted(missing_in_reqs))}
                </ul>
            </div>
            
            <div id="format-errors" class="section">
                <h2>🚨 Formatting Rule Violations</h2>
                <p><em>Rule: Must end with a comma, contain NO spaces, and NO newlines.</em></p>
                <ul>
                    {format_errors_html}
                </ul>
            </div>
        </div>
    </body>
    </html>
    """
    
    with open(output_file, 'w', encoding='utf-8') as f:
        f.write(html_content)
        
    logging.info(f"Report successfully generated: {os.path.abspath(output_file)}")

def main():
    parser = argparse.ArgumentParser(description="Generate 2-way UT Traceability Matrix and check formatting.")
    parser.add_argument("req_file", help="Path to the Requirements Excel file")
    # Using nargs='+' allows the user to pass one or multiple UT files separated by space
    parser.add_argument("ut_files", nargs='+', help="Paths to one or more UT Excel files")
    parser.add_argument("-o", "--output", default="ut_traceability_report.html", help="Output HTML report file name")
    
    args = parser.parse_args()
    
    if not os.path.exists(args.req_file):
        logging.error(f"Requirements file not found: {args.req_file}")
        return
        
    for ut_file in args.ut_files:
        if not os.path.exists(ut_file):
            logging.error(f"UT file not found: {ut_file}")
            return
            
    req_ids = extract_reqs_from_req_file(args.req_file)
    ut_reqs, format_errors = extract_reqs_from_ut_files(args.ut_files)
    
    if not req_ids and not ut_reqs:
        logging.error("No requirements found in both files. Exiting.")
        return
        
    generate_html_report(req_ids, ut_reqs, format_errors, args.output)

if __name__ == "__main__":
    main()